"""Funções comuns: config, log, navegador, consulta e planilha."""
import json
import logging
import re
import subprocess
import time
from datetime import datetime, date
from pathlib import Path

import openpyxl
from playwright.sync_api import TimeoutError as PWTimeout

BASE = Path(__file__).resolve().parent
log = logging.getLogger("comprovei")
STATUS_FINAIS = {"COMPLETA", "INCOMPLETA", "NÃO ENCONTRADA"}


class ConsultaErro(Exception):
    """Falha técnica (pode ser repetida)."""


class DadosInvalidos(ConsultaErro):
    """Dado de entrada inválido (não adianta repetir)."""


# ---------------------------------------------------------------- config / log
def carregar_config(nome="config.json"):
    return json.loads((BASE / nome).read_text(encoding="utf-8"))


def configurar_log():
    pasta = BASE / "logs"
    pasta.mkdir(exist_ok=True)
    arq = pasta / f"automacao_{datetime.now():%Y%m%d}.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        handlers=[logging.FileHandler(arq, encoding="utf-8"), logging.StreamHandler()],
    )
    return log


# ------------------------------------------------------------------- navegador
def abrir_navegador(p, cfg):
    """Retorna (context, page). Modo 'cdp' (padrão): o Chrome fica aberto
    após o script terminar, mantendo a sessão. Modo 'launch': o Playwright
    abre o Chrome e ele fecha quando o script termina."""
    chrome, perfil = cfg["chrome_caminho"], cfg["perfil_dir"]
    if not Path(chrome).exists():
        raise FileNotFoundError(f"Chrome não encontrado em: {chrome}")
    Path(perfil).mkdir(parents=True, exist_ok=True)

    if cfg["modo_navegador"] == "launch":
        ctx = p.chromium.launch_persistent_context(
            perfil, executable_path=chrome, headless=False,
            no_viewport=True, args=["--start-maximized"])
        return ctx, (ctx.pages[0] if ctx.pages else ctx.new_page())

    porta = cfg["porta_depuracao"]
    url = f"http://127.0.0.1:{porta}"
    browser = None
    try:
        browser = p.chromium.connect_over_cdp(url)
    except Exception:
        log.info("Abrindo o Chrome com perfil separado (porta %s)...", porta)
        subprocess.Popen([
            chrome, f"--remote-debugging-port={porta}",
            f"--user-data-dir={perfil}", "--no-first-run",
            "--no-default-browser-check",
            cfg.get("url_sistema") or "about:blank"])
        for _ in range(30):
            time.sleep(1)
            try:
                browser = p.chromium.connect_over_cdp(url)
                break
            except Exception:
                pass
    if browser is None:
        raise RuntimeError(
            "Não consegui conectar ao Chrome. Feche TODAS as janelas do Chrome "
            "que usam este perfil e tente de novo.")
    ctx = browser.contexts[0]
    return ctx, (ctx.pages[0] if ctx.pages else ctx.new_page())


# --------------------------------------------------------------------- consulta
def formatar_data(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    s = str(v).strip()
    for texto, fmt in ((s, "%d/%m/%Y"), (s[:10], "%Y-%m-%d"), (s, "%d/%m/%y")):
        try:
            return datetime.strptime(texto, fmt).strftime("%d/%m/%Y")
        except ValueError:
            pass
    raise DadosInvalidos(f"Data inválida: {v!r} (use DD/MM/AAAA)")


def norm_carga(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _campo(page, cfg, chave, rotulo):
    css = cfg["seletores_css"].get(chave)
    if css:
        return page.locator(css).first
    # 1) label associado  2) placeholder
    rx = re.compile(rf"^\s*{re.escape(rotulo)}\s*\*?\s*(Max.*)?$", re.I)
    loc = page.get_by_label(rx)
    if loc.count() >= 1:
        return loc.first
    loc = page.get_by_placeholder(rx)
    if loc.count() >= 1:
        return loc.first
    raise ConsultaErro(
        f"Campo '{rotulo}' não localizado. Rode inspecionar.py e preencha "
        f"seletores_css.{chave} no config.json.")


def _preencher(campo, texto):
    """fill() não depende de clique (o datepicker pode cobrir campos vizinhos)."""
    alvo = re.sub(r"\D", "", texto)
    campo.fill(texto)
    campo.evaluate(
        "e => { for (const t of ['input','keyup','change']) "
        "e.dispatchEvent(new Event(t,{bubbles:true})); e.blur(); }")
    if re.sub(r"\D", "", campo.input_value() or "") != alvo:
        # fallback: digitação, como no primeiro teste
        campo.click(force=True)
        campo.press("Control+A")
        campo.press_sequentially(texto, delay=30)
        campo.press("Tab")
        if re.sub(r"\D", "", campo.input_value() or "") != alvo:
            log.warning("Campo ficou com '%s' em vez de '%s'", campo.input_value(), texto)


def _url_consulta(cfg):
    url = cfg.get("url_consulta") or ""
    n = cfg.get("linhas_por_pagina", 250)
    if not url:
        return ""
    if re.search(r"limit=\d+", url):
        return re.sub(r"limit=\d+", f"limit={n}", url)
    return url + ("&" if "?" in url else "?") + f"limit={n}"


def _garantir_tela(page, cfg):
    url = _url_consulta(cfg)
    if url:
        page.goto(url, wait_until="domcontentloaded")   # estado limpo a cada consulta
    t = cfg["textos"]["botao_filtros"]
    btn = page.get_by_role("button", name=re.compile(rf"^\s*{t}\s*$")).or_(
        page.get_by_text(re.compile(rf"^\s*{t}\s*$"))).filter(visible=True)
    try:
        btn.first.wait_for(state="visible", timeout=cfg["timeout_ms"])
    except PWTimeout:
        raise ConsultaErro(
            "Tela ROTAS/CONSULTA não encontrada (sessão expirada? faça login "
            "no Chrome ou confira url_consulta).")
    return btn.first


def _aguardar_modal_fechado(page):
    try:
        page.locator(".modal-backdrop").first.wait_for(state="detached", timeout=3000)
    except PWTimeout:
        pass


def _clicar_filtrar(page, cfg, campo_rota):
    """Clica em Filtrar pelo id (#filterButton); se algo cobrir o botão, clica via JS."""
    css = cfg["seletores_css"].get("botao_filtrar") or "#filterButton"
    alvo = page.locator(css).filter(visible=True).first
    try:
        alvo.click(timeout=3000)
        return
    except PWTimeout:
        pass
    try:
        alvo.evaluate("e => e.click()", timeout=3000)
        return
    except Exception:
        pass
    log.warning("Botão Filtrar (%s) não respondeu; enviei Enter no campo Rota.", css)
    campo_rota.press("Enter")


def _submeter(page, cfg, campo_rota):
    """Envia o filtro e ESPERA a página recarregar (evita ler a tabela antiga)."""
    try:
        with page.expect_navigation(wait_until="domcontentloaded",
                                    timeout=cfg.get("timeout_navegacao_ms", 10000)):
            _clicar_filtrar(page, cfg, campo_rota)
    except PWTimeout:
        log.warning("Filtro enviado, mas a página não recarregou em tempo; seguindo.")


_JS_SUBMIT = """(d) => {
  const get = (sel) => { const e = document.querySelector(sel);
    if (!e) throw new Error('campo não encontrado: ' + sel); return e; };
  get(d.rota).value = d.numero;
  get(d.de).value = d.de_v;
  get(d.ate).value = d.ate_v;
  const btn = get(d.botao);
  const f = btn.form;
  if (!f) throw new Error('botão sem formulário: ' + d.botao);
  if (f.requestSubmit) f.requestSubmit(btn); else f.submit();
}"""


def _aplicar_filtro_js(page, cfg, rota_txt, de_s, ate_s):
    """Preenche os campos pelo id e envia o formulário, sem abrir o painel nem clicar."""
    sel = cfg["seletores_css"]
    ids = {
        "rota": sel.get("campo_rota") or "#RoutesNumber",
        "de": sel.get("campo_de") or "#RoutesRouteDate1",
        "ate": sel.get("campo_ate") or "#RoutesRouteDate2",
        "botao": sel.get("botao_filtrar") or "#filterButton",
    }
    url = _url_consulta(cfg)
    if url:
        page.goto(url, wait_until="load")
    try:
        page.wait_for_selector(ids["rota"], state="attached", timeout=cfg["timeout_ms"])
    except PWTimeout:
        raise ConsultaErro("Tela de consulta não encontrada (sessão expirada? rode login.py).")
    dados = {**ids, "numero": rota_txt, "de_v": de_s, "ate_v": ate_s}
    try:
        with page.expect_navigation(wait_until="load",
                                    timeout=cfg.get("timeout_navegacao_ms", 15000)):
            page.evaluate(_JS_SUBMIT, dados)
    except PWTimeout:
        raise ConsultaErro("O filtro foi enviado, mas a página não recarregou.")
    try:
        page.wait_for_load_state("networkidle", timeout=cfg.get("espera_rede_ms", 3000))
    except PWTimeout:
        pass
    page.wait_for_timeout(cfg["espera_pos_filtro_ms"])


def _indice_coluna(page, nome):
    hs = page.locator("th, [role='columnheader']")
    n = hs.count()
    for i in range(n):
        if nome.lower() in " ".join(hs.nth(i).inner_text().split()).lower():
            return i, n
    return None, n


def consultar_carga(page, cfg, carga, de, ate, pasta_prints=None):
    """Retorna (x, y) ou None se não encontrada."""
    t = cfg["textos"]
    to = cfg["timeout_ms"]
    de_s, ate_s = formatar_data(de), formatar_data(ate)
    d1 = datetime.strptime(de_s, "%d/%m/%Y")
    d2 = datetime.strptime(ate_s, "%d/%m/%Y")
    if d2 < d1:
        raise DadosInvalidos("Data final anterior à data inicial")
    if (d2 - d1).days > 31:
        raise DadosInvalidos("Período maior que 31 dias (limite do sistema)")
    if not carga:
        raise DadosInvalidos("Carga/Rota vazia")

    aplicar_filtro(page, cfg, carga, de_s, ate_s)

    padrao = re.compile(rf"(?<!\d){re.escape(carga)}(?!\d)")
    css_linhas = cfg["seletores_css"].get("linhas") or \
        "tr:not(:has(tr)), [role='row']:not(:has([role='row']))"
    xy = re.compile(r"(?<![\d/])\d+\s*/\s*\d+(?![\d/])")
    linhas = page.locator(css_linhas).filter(has_text=padrao).filter(has_text=xy)
    limite = time.time() + cfg["espera_resultado_ms"] / 1000
    while linhas.count() == 0 and time.time() < limite:
        time.sleep(0.5)
    if pasta_prints:
        page.screenshot(path=str(pasta_prints / f"{carga}_2_resultado.png"))
    if linhas.count() == 0:
        return None
    if linhas.count() > 1:
        log.warning("Carga %s: %d linhas encontradas; usando a primeira.", carga, linhas.count())

    linha = linhas.first
    idx, n_cab = _indice_coluna(page, t["coluna_resultado"])
    celulas = linha.locator("td, [role='cell'], [role='gridcell']")
    if idx is not None and celulas.count() == n_cab:
        texto = celulas.nth(idx).inner_text()
    else:
        log.warning("Colunas não alinham com o cabeçalho; lendo a linha inteira.")
        texto = linha.inner_text()
    m = re.search(r"(?<![\d/])(\d+)\s*/\s*(\d+)(?![\d/])", texto)
    if not m:
        raise ConsultaErro(f"Linha encontrada, mas sem valor x/y em '{t['coluna_resultado']}': {texto!r}")
    return int(m.group(1)), int(m.group(2))


def classificar(res):
    """Retorna (texto_resultado, status, observacao)."""
    if res is None:
        return "", "NÃO ENCONTRADA", "Carga não localizada no período informado"
    x, y = res
    txt = f"{x}/{y}"
    if y == 0:
        raise ConsultaErro(f"Total de documentos igual a zero ({txt})")
    if x > y:
        raise ConsultaErro(f"Valor inconsistente ({txt})")
    if x == y:
        return txt, "COMPLETA", "Todos os documentos estão finalizados"
    return txt, "INCOMPLETA", f"Existem {x} documentos finalizados de {y} esperados (faltam {y - x})"


def executar_consulta(page, cfg, carga, de, ate, pasta_prints=None):
    """Consulta com tentativas. Nunca levanta exceção: devolve ERRO."""
    ultimo = ""
    for tentativa in range(1, cfg["tentativas"] + 1):
        try:
            res = consultar_carga(page, cfg, carga, de, ate, pasta_prints)
            if res is None and cfg.get("confirmar_nao_encontrada", True):
                log.info("Carga %s: não encontrada; confirmando com nova consulta...", carga)
                res = consultar_carga(page, cfg, carga, de, ate, pasta_prints)
            return classificar(res)
        except DadosInvalidos as e:
            return "", "ERRO", str(e)
        except Exception as e:  # ConsultaErro, timeout, etc.
            ultimo = f"{type(e).__name__}: {' | '.join(l.strip() for l in str(e).splitlines()[:4])}"
            log.warning("Carga %s, tentativa %d falhou: %s", carga, tentativa, ultimo)
            try:
                pasta = BASE / "prints"
                pasta.mkdir(exist_ok=True)
                page.screenshot(path=str(pasta / f"erro_{carga}_{datetime.now():%H%M%S}.png"))
            except Exception:
                pass
            time.sleep(1)
    return "", "ERRO", ultimo


# --------------------------------------------------------------------- planilha
def abrir_planilha(cfg):
    caminho = Path(cfg["planilha"])
    if not caminho.is_absolute():
        caminho = BASE / caminho
    if not caminho.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {caminho}")
    if caminho.with_name(f".~lock.{caminho.name}#").exists():
        log.warning("A planilha parece aberta no LibreOffice. Feche-a antes de processar.")
    wb = openpyxl.load_workbook(caminho)
    ws = wb[cfg["aba"]] if cfg.get("aba") else wb.active
    cab = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    col = {}
    for chave, nome in cfg["colunas"].items():
        if nome in cab:
            col[chave] = cab[nome]
        elif chave in ("resultado", "status", "observacao"):
            nova = ws.max_column + 1
            ws.cell(1, nova, nome)
            col[chave] = nova
        else:
            raise ValueError(f"Coluna '{nome}' não encontrada. Colunas existentes: {list(cab)}")
    return caminho, wb, ws, col


# ------------------------------------------------------------ consulta em lote
_XY = re.compile(r"(?<![\d/])(\d+)\s*/\s*(\d+)(?![\d/])")
_CSS_LINHAS = "tr:not(:has(tr)), [role='row']:not(:has([role='row']))"


def aplicar_filtro(page, cfg, rota_txt, de_s, ate_s):
    """Aplica o filtro. Modo 'js' (padrão): envia o formulário direto. Modo 'tela': usa o painel."""
    if cfg.get("modo_filtro", "js") == "js":
        return _aplicar_filtro_js(page, cfg, rota_txt, de_s, ate_s)
    t = cfg["textos"]
    to = cfg["timeout_ms"]
    btn_filtros = _garantir_tela(page, cfg)
    try:
        aberto = _campo(page, cfg, "campo_rota", t["rotulo_rota"]).is_visible()
    except ConsultaErro:
        aberto = False
    if not aberto:
        btn_filtros.click()
    campo_rota = _campo(page, cfg, "campo_rota", t["rotulo_rota"])
    campo_rota.wait_for(state="visible", timeout=to)
    _preencher(_campo(page, cfg, "campo_de", t["rotulo_de"]), de_s)
    _preencher(_campo(page, cfg, "campo_ate", t["rotulo_ate"]), ate_s)
    _preencher(campo_rota, rota_txt)
    if re.sub(r"\D", "", campo_rota.input_value() or "") != re.sub(r"\D", "", rota_txt):
        raise ConsultaErro("Campo Rota não aceitou todas as cargas (limite de tamanho?)")
    _submeter(page, cfg, campo_rota)
    try:
        campo_rota.wait_for(state="hidden", timeout=to)
    except PWTimeout:
        raise ConsultaErro("O painel de filtros não fechou após 'Filtrar'")
    _aguardar_modal_fechado(page)
    try:
        page.wait_for_load_state("networkidle", timeout=cfg.get("espera_rede_ms", 3000))
    except PWTimeout:
        pass
    page.wait_for_timeout(cfg["espera_pos_filtro_ms"])


def _linhas_atuais(page):
    try:
        tx = page.get_by_text(re.compile(r"Linhas:\s*\d+")).first.inner_text(timeout=2000)
        return int(re.search(r"\d+", tx.split(":")[1]).group())
    except Exception:
        return None


def _definir_linhas(page, cfg, n):
    if (_linhas_atuais(page) or 0) >= n:
        return
    try:
        page.get_by_text(re.compile(r"Linhas:\s*\d+")).first.click(timeout=3000)
        page.get_by_text(str(n), exact=True).filter(visible=True).first.click(timeout=3000)
        try:
            page.wait_for_load_state("networkidle", timeout=cfg.get("espera_rede_ms", 3000))
        except PWTimeout:
            pass
        page.wait_for_timeout(cfg["espera_pos_filtro_ms"])
    except Exception as e:
        log.warning("Não consegui mudar 'Linhas' para %s (%s). Seguindo com %s.",
                    n, str(e).splitlines()[0], _linhas_atuais(page))


def _ler_textos_linhas(page, cfg):
    css = cfg["seletores_css"].get("linhas") or _CSS_LINHAS
    loc = page.locator(css).filter(has_text=_XY)
    limite = time.time() + cfg["espera_resultado_ms"] / 1000
    while loc.count() == 0 and time.time() < limite:
        time.sleep(0.4)
    anterior = -1
    for _ in range(6):                      # espera a tabela estabilizar
        n = loc.count()
        if n == anterior:
            break
        anterior = n
        time.sleep(0.3)
    return loc.all_inner_texts()


def consultar_lote(page, cfg, cargas, de, ate, pasta_prints=None, _segunda=False):
    """Consulta várias cargas de uma vez (Rota separada por vírgula).
    Retorna dict: carga -> (x, y) | None (comprovadamente não encontrada).
    Cargas AUSENTES do dict ficaram indeterminadas (o chamador consulta uma a uma)."""
    de_s, ate_s = formatar_data(de), formatar_data(ate)
    d1 = datetime.strptime(de_s, "%d/%m/%Y")
    d2 = datetime.strptime(ate_s, "%d/%m/%Y")
    if d2 < d1:
        raise DadosInvalidos("Data final anterior à data inicial")
    if (d2 - d1).days > 31:
        raise DadosInvalidos("Período maior que 31 dias (limite do sistema)")

    aplicar_filtro(page, cfg, ",".join(cargas), de_s, ate_s)
    _definir_linhas(page, cfg, cfg.get("linhas_por_pagina", 250))
    textos = _ler_textos_linhas(page, cfg)
    if pasta_prints:
        page.screenshot(path=str(pasta_prints / f"lote_{datetime.now():%H%M%S}.png"))

    alvo = set(cargas)
    res, ambiguo = {}, False
    for tx in textos:
        m = _XY.search(tx)
        achadas = set(re.findall(r"(?<!\d)\d{4,}(?!\d)", tx)) & alvo
        if not m or not achadas:
            continue
        if len(achadas) > 1:
            ambiguo = True
            continue
        res[achadas.pop()] = (int(m.group(1)), int(m.group(2)))

    limite = _linhas_atuais(page) or cfg.get("linhas_por_pagina", 250)
    truncado = len(textos) >= limite
    log.info("Lote%s: %d cargas pedidas, %d linhas lidas, %d encontradas (linhas/página=%s)",
             " (confirmação)" if _segunda else "", len(cargas), len(textos), len(res), limite)
    if not textos:
        # tabela vazia: na confirmação significa "não existe"; na 1ª passada é duvidoso
        return {c: None for c in cargas} if _segunda else {}
    if not res:
        log.warning("Lote: há linhas na tabela mas nenhuma é das cargas pedidas "
                    "(filtro possivelmente não aplicado). Consultando individualmente.")
        return {}
    if truncado or ambiguo:
        return res
    faltam = [c for c in cargas if c not in res]
    if faltam:
        if not _segunda and cfg.get("confirmar_nao_encontrada", True):
            res.update(consultar_lote(page, cfg, faltam, de, ate, pasta_prints, _segunda=True))
        else:
            for c in faltam:
                res[c] = None
    return res
